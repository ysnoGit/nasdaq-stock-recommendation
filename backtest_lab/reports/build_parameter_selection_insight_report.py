from __future__ import annotations

from pathlib import Path
import sys

from docx import Document
from docx.shared import Inches, Pt

sys.path.append(str(Path(__file__).resolve().parents[2]))

from backtest_lab.reports.build_backtest_plan_flow_report import (
    INK,
    MUTED,
    PALE_GOLD,
    PALE_GREEN,
    add_bullet,
    add_callout,
    add_heading,
    add_table,
    configure_document,
)


REPORT_DIR = Path(__file__).resolve().parent
OUTPUT = REPORT_DIR / "backtest_parameter_selection_insight_report.docx"


def build_report() -> Path:
    doc = Document()
    configure_document(doc)
    doc.sections[0].top_margin = Inches(0.7)
    doc.sections[0].bottom_margin = Inches(0.7)
    doc.styles["Normal"].font.size = Pt(10)
    doc.styles["List Bullet"].font.size = Pt(10)
    doc.sections[0].header.paragraphs[0].text = (
        "NASDAQ Stock Recommendation | Parameter Selection Decision"
    )

    title = doc.add_paragraph()
    run = title.add_run("PARAMETER SET SELECTION INSIGHT")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = INK
    subtitle = doc.add_paragraph()
    subtitle.add_run(
        "Decision synthesis from the corrected timing flow, screening-yield comparison, "
        "and fixed-horizon performance comparison"
    ).font.color.rgb = MUTED

    add_callout(
        doc,
        "Recommended default",
        "Use parameter set 17 (ag2_qg2_ay2_qc4_vr2_vd2) as the primary balanced "
        "configuration. It combines strong cross-horizon returns, favorable median returns "
        "and win rates, and adequate completed-sample coverage in both A-F and A-H.",
        PALE_GREEN,
    )

    add_heading(doc, "Decision Summary")
    add_table(
        doc,
        ["Role", "Parameter set", "Configuration", "Why it belongs"],
        [
            [
                "Primary balanced default",
                "17",
                "A/Q growth 2%/2%; 2 annual years; 4 quarters; 2x volume; 2 surge days",
                "Best overall quality consistency across means, medians, win rates, horizons, and screens.",
            ],
            [
                "Return-ranked challenger",
                "113",
                "A/Q growth 3%/2%; 2 annual years; 4 quarters; 2x volume; 2 surge days",
                "Best A-F combined average-return rank and strong A-H results, with slightly less coverage.",
            ],
            [
                "Broad discovery baseline",
                "1",
                "A/Q growth 2%/2%; 2 annual years; 2 quarters; 2x volume; 2 surge days",
                "Largest useful candidate pool and strongest coverage; weaker short-term selectivity.",
            ],
        ],
        [1.25, 0.75, 2.45, 2.15],
    )

    add_heading(doc, "Why Set 17 Is The Best Default")
    add_bullet(
        doc,
        "It ranks first on the report's broader quality assessment for both A-F and A-H "
        "when average returns, median returns, and win rates are considered together.",
    )
    add_bullet(
        doc,
        "Its completed samples remain meaningful: A-F has 56/53/43 observations at "
        "6 months/1 year/2 years, while A-H has 29/28/22.",
    )
    add_bullet(
        doc,
        "Its four-quarter requirement improves durability without using the stricter "
        "three-year annual-history requirement that sharply reduces screening yield.",
    )
    add_bullet(
        doc,
        "Its 2x volume threshold avoids the severe sample collapse observed at 3x-5x thresholds.",
    )

    add_table(
        doc,
        ["Screen", "N 6m/1y/2y", "Average return 6m/1y/2y", "Median return 6m/1y/2y", "Win rate 6m/1y/2y"],
        [
            ["A-F", "56 / 53 / 43", "2.75% / 8.52% / 9.51%", "2.60% / 6.96% / 5.89%", "57.14% / 60.38% / 55.81%"],
            ["A-H", "29 / 28 / 22", "4.98% / 10.43% / 1.56%", "1.75% / 8.06% / -1.52%", "55.17% / 75.00% / 50.00%"],
        ],
        [0.7, 1.0, 1.7, 1.7, 1.5],
    )

    add_heading(doc, "Set 17 Versus Set 113")
    add_table(
        doc,
        ["Question", "Set 17", "Set 113", "Interpretation"],
        [
            ["Annual growth threshold", "2%", "3%", "Set 113 demands slightly stronger annual growth."],
            ["A-F average returns", "2.75 / 8.52 / 9.51%", "2.68 / 8.35 / 10.22%", "Very similar; set 113 leads only at 2 years."],
            ["A-F completed N", "56 / 53 / 43", "54 / 51 / 41", "Set 17 has modestly better coverage."],
            ["A-H average returns", "4.98 / 10.43 / 1.56%", "5.05 / 10.11 / 2.24%", "Set 113 slightly improves 6m and 2y averages."],
            ["Robustness", "Stronger medians/win-rate blend", "Best average-return rank", "Choose 17 unless maximizing average-return rank is the only objective."],
        ],
        [1.25, 1.45, 1.45, 2.25],
    )
    add_callout(
        doc,
        "Practical conclusion",
        "The evidence does not justify treating set 113 as decisively superior to set 17. "
        "Their observations substantially overlap, and the differences are small. Set 17's "
        "slightly broader coverage and stronger median/win-rate profile make it the safer default.",
    )

    add_heading(doc, "How To Use A-F And A-H")
    add_bullet(
        doc,
        "Use A-F as the primary product candidate list. It has materially larger completed "
        "samples and more credible two-year evidence.",
    )
    add_bullet(
        doc,
        "Use A-H as an additional confidence or timing label, not as proof of superior "
        "long-term performance.",
    )
    add_bullet(
        doc,
        "A-H shows attractive 6-month and 1-year results for set 17, but its 2-year average "
        "falls to 1.56%, its 2-year median is -1.52%, and its 2-year win rate is 50%.",
    )
    add_callout(
        doc,
        "Important interpretation",
        "Weekly G/H confirmation appears useful for short- and medium-term selectivity, but "
        "the current evidence does not support using A-H as a stronger two-year holding signal.",
        PALE_GOLD,
    )

    add_heading(doc, "Parameter-Level Insights")
    add_table(
        doc,
        ["Parameter choice", "Observed implication", "Recommendation"],
        [
            ["Volume ratio threshold", "Increasing from 2x toward 5x sharply reduces both A-F and A-H yield.", "Keep 2x as the default; treat 3x+ as optional high-conviction filters."],
            ["Annual periods", "Three years substantially reduces candidate coverage.", "Use 2 years by default until more history increases sample size."],
            ["Quarterly periods", "Four quarters improves the strongest balanced configurations.", "Use 4 quarters for the primary strategy."],
            ["Growth thresholds", "Moving from 2% to 3% changes results modestly compared with volume and period choices.", "Prefer 2% default; keep 3% annual growth as a challenger."],
            ["Volume surge days", "Two days generally preserves more evidence than three days.", "Use 2 days as the default."],
        ],
        [1.3, 3.1, 2.0],
    )

    add_heading(doc, "Recommended Product Configuration")
    add_table(
        doc,
        ["Setting", "Recommended value"],
        [
            ["Primary parameter set", "17: ag2_qg2_ay2_qc4_vr2_vd2"],
            ["Annual growth threshold", "2%"],
            ["Quarterly growth threshold", "2%"],
            ["Annual periods", "2"],
            ["Quarterly periods", "4"],
            ["Volume ratio threshold", "2x"],
            ["Volume surge minimum days", "2"],
            ["Primary result list", "A-F"],
            ["Additional label", "A-H confirmed"],
            ["Challenger monitored beside default", "113: annual growth threshold increased to 3%"],
        ],
        [2.5, 3.8],
    )

    add_heading(doc, "What Would Change The Recommendation")
    add_bullet(doc, "A materially larger A-H two-year sample that produces positive median returns and a win rate clearly above 50%.")
    add_bullet(doc, "Walk-forward or out-of-sample testing showing set 113 consistently outperforming set 17.")
    add_bullet(doc, "Fundamental filing-availability dates that remove the remaining look-ahead-bias risk.")
    add_bullet(doc, "Portfolio simulations incorporating repeated entries, liquidity, transaction costs, position sizing, and overlapping capital.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_report())
